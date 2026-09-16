"""文本型 PDF 双链路路由（Issue #61）：
MASK=实体定位→整页栅格化真打码（与扫描件同构）；
替换模式=PDF→docx→替换→PDF 回转，转换失败回退原位替换。
"""

import asyncio
import os

import fitz
import pytest

from app.core.config import settings
from app.models.entity_schemas import BoundingBox, Entity
from app.models.redaction_schemas import RedactionConfig
from app.services.redactor import Redactor


@pytest.fixture()
def _dirs(tmp_path, monkeypatch):
    up = tmp_path / "uploads"; out = tmp_path / "outputs"
    up.mkdir(); out.mkdir()
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(up))
    monkeypatch.setattr(settings, "OUTPUT_DIR", str(out))
    return up, out


def _make_pdf(path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 130), "委托人：陈明飞，身份证号 110101199003078515", fontsize=12, fontname="china-s")
    doc.save(str(path)); doc.close()


def _entities():
    return [
        Entity(id="e1", text="陈明飞", type="PERSON", start=0, end=3, page=1, selected=True),
        Entity(id="e2", text="110101199003078515", type="ID_CARD", start=0, end=18, page=1, selected=True),
    ]


def _page_has_images(pdf_path):
    doc = fitz.open(pdf_path)
    try:
        imgs = doc.load_page(0).get_images(full=True)
    finally:
        doc.close()
    return len(imgs) > 0


@pytest.mark.asyncio
async def test_pdf_mask_rasterizes_and_mosaics(_dirs):
    """MASK 模式：文本型 PDF 栅格化为图像页，原文从文本层消失。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=_entities(), bounding_boxes=[],
        config=RedactionConfig(replacement_mode="mask"),
    )
    assert result["redacted_count"] >= 2
    doc = fitz.open(result["output_path"])
    try:
        text = doc.load_page(0).get_text()
    finally:
        doc.close()
    assert "陈明飞" not in text and "110101199003078515" not in text
    assert _page_has_images(result["output_path"]), "MASK 产物应含栅格化图像页"
    assert result["residual_entities"] == []


@pytest.mark.asyncio
async def test_pdf_mask_missing_entity_falls_back_to_text_mask(_dirs):
    """实体在文本层定位失败（跨行等）：整份回退文本掩码链路（栅格化会把
    漏网原文留在可读像素里，零容忍），未定位实体进 residual_entities。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    ghost = Entity(id="e3", text="不存在的实体文本XYZ", type="PERSON", start=0, end=3, page=1, selected=True)
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=_entities() + [ghost], bounding_boxes=[],
        config=RedactionConfig(replacement_mode="mask"),
    )
    doc = fitz.open(result["output_path"])
    try:
        text = doc.load_page(0).get_text()
        has_imgs = len(doc.load_page(0).get_images(full=True)) > 0
    finally:
        doc.close()
    assert "陈明飞" not in text.replace(" ", ""), "回退文本链路后可定位实体原文必须删除"
    assert not has_imgs, "存在漏定位实体时禁止栅格化（明文像素泄露）"
    # 残留以详细版（含替换去向）或纯文本形态透出均可，关键是不静默
    assert any("不存在的实体文本XYZ" in r for r in result["residual_entities"])


@pytest.mark.asyncio
async def test_pdf_replacement_routes_via_docx_roundtrip(_dirs):
    """替换模式：走 pdf→docx→替换→PDF 回转；LibreOffice 缺失时自动回退
    原位替换——两种产物都必须满足『原文消失+替换词落盘』契约。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    pools = {"PERSON": {"words": ["王某"], "strategy": "numbered", "custom_map": {}}}
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=_entities(), bounding_boxes=[],
        config=RedactionConfig(replacement_mode="pseudonym", word_pools=pools),
    )
    text = "".join(p.get_text() for p in fitz.open(result["output_path"]))
    assert "陈明飞" not in text.replace(" ", "")
    assert "110101199003078515" not in text.replace(" ", "")
    assert "王某" in text.replace(" ", ""), "替换词未落盘"
    assert result["redacted_count"] == 2


@pytest.mark.asyncio
async def test_pdf_replacement_prefers_docx_roundtrip(_dirs, monkeypatch):
    """路由断言：替换模式必须优先尝试 docx 回转链路（而非静默走原位）。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    called = {}

    def _fake_pdf2docx(src, wd):
        from docx import Document as _Doc
        import shutil as _sh
        fake_docx = os.path.join(wd, "source.docx")
        d = _Doc()
        for e in _entities():
            d.add_paragraph(e.text)
        d.save(fake_docx)
        called["pdf2docx"] = True
        return fake_docx

    monkeypatch.setattr(Redactor, "_pdf_to_docx", staticmethod(_fake_pdf2docx))
    async def _fake_docx2pdf(docx, out):
        called["docx2pdf"] = True
        import shutil; shutil.copy(str(src), out); return True
    monkeypatch.setattr(Redactor, "_docx_to_pdf", staticmethod(_fake_docx2pdf))
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=_entities(), bounding_boxes=[],
        config=RedactionConfig(replacement_mode="structured"),
    )
    assert called.get("pdf2docx") and called.get("docx2pdf"), "替换模式未走 docx 回转链路"


@pytest.mark.asyncio
async def test_pdf_mask_with_user_boxes_also_masks_entities(_dirs, monkeypatch):
    """拉框 + MASK：用户框与实体定位框**合并**栅格化（评审 I2）——只吃
    手拉框会在「选了实体+残留旧框」时产出零打码的栅格化成品。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    called = {}
    _orig_locate = Redactor._entities_to_norm_boxes  # py3.10+ 类属性已是裸函数

    def _spy(file_path, ents):
        called["located"] = True
        return _orig_locate(file_path, ents)

    monkeypatch.setattr(Redactor, "_entities_to_norm_boxes", staticmethod(_spy))
    boxes = [BoundingBox(id="b1", x=0.1, y=0.1, width=0.3, height=0.05, page=1, type="PERSON", selected=True)]
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=_entities(), bounding_boxes=boxes,
        config=RedactionConfig(replacement_mode="mask"),
    )
    assert called.get("located"), "拉框路径的 MASK 也应做实体定位并合并"
    assert result["redacted_count"] == 3, "1 用户框 + 2 实体框"
    doc = fitz.open(result["output_path"])
    try:
        text = "".join(p.get_text() for p in doc)
    finally:
        doc.close()
    assert "陈明飞" not in text.replace(" ", "") and "110101199003078515" not in text


@pytest.mark.asyncio
async def test_pdf_replacement_with_user_boxes_routes_to_image_pipeline(_dirs):
    """拉框 + 替换模式：拉框是视觉标注信号，走图像管线（既有行为）。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    boxes = [BoundingBox(id="b1", x=0.1, y=0.1, width=0.3, height=0.05, page=1, type="PERSON", selected=True)]
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=_entities(), bounding_boxes=boxes,
        config=RedactionConfig(replacement_mode="structured"),
    )
    assert result["redacted_count"] == 1
    assert _page_has_images(result["output_path"]), "拉框路径产物应栅格化"


def test_entities_to_norm_boxes_coordinates(_dirs):
    """实体定位框为归一化坐标且落在页面范围内。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    boxes, missed = Redactor._entities_to_norm_boxes(str(src), _entities())
    assert not missed
    assert len(boxes) == 2
    for b in boxes:
        assert 0 <= b.x < 1 and 0 <= b.y < 1
        assert 0 < b.width <= 1 and 0 < b.height <= 1
        assert b.selected and b.page == 1


def test_entities_to_norm_boxes_searches_all_pages(_dirs):
    """同一实体文本在多页出现时必须全部打框（评审 C1）——NER 漏检重复
    出现时，只按登记页打码会让其他页原文以可读像素残留。"""
    up, _ = _dirs
    src = up / "t.pdf"
    doc = fitz.open()
    for _ in range(3):
        doc.new_page()
    for i in range(3):
        doc[i].insert_text((72, 130), "嫌疑人陈明飞到案", fontsize=12, fontname="china-s")
    doc.save(str(src)); doc.close()
    # 实体只登记第 1 页（模拟 NER 漏检后两页的重复出现）
    ents = [Entity(id="e1", text="陈明飞", type="PERSON", start=0, end=3, page=1, selected=True)]
    boxes, missed = Redactor._entities_to_norm_boxes(str(src), ents)
    assert not missed
    assert {b.page for b in boxes} == {1, 2, 3}, "三页同名文本必须都有框"


@pytest.mark.asyncio
async def test_pdf_mask_masks_same_text_on_all_pages(_dirs):
    """端到端（评审 C1）：实体只登记第 1 页，第 2 页同名原文也必须被打码，
    产物任何页都不得残留原文文本层。"""
    up, _ = _dirs
    src = up / "t2.pdf"
    doc = fitz.open()
    for _ in range(2):
        doc.new_page()
    for i in range(2):
        doc[i].insert_text((72, 130), "委托人：陈明飞", fontsize=12, fontname="china-s")
    doc.save(str(src)); doc.close()
    ents = [Entity(id="e1", text="陈明飞", type="PERSON", start=0, end=3, page=1, selected=True)]
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=ents, bounding_boxes=[],
        config=RedactionConfig(replacement_mode="mask"),
    )
    out = fitz.open(result["output_path"])
    try:
        for i in range(out.page_count):
            assert "陈明飞" not in out[i].get_text(), f"第{i+1}页文本层残留原文"
            assert len(out[i].get_images(full=True)) > 0, f"第{i+1}页应栅格化"
    finally:
        out.close()
    assert result["redacted_count"] == 2, "两页同名各一框"
    assert result["residual_entities"] == []


@pytest.mark.asyncio
async def test_pdf_replacement_zero_entities_short_circuits(_dirs, monkeypatch):
    """零实体（评审 I3）：直接原样拷贝，不跑 docx 回转（避免无意义的有损
    重排），输出与输入逐字节一致。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    called = {}
    monkeypatch.setattr(
        Redactor, "_pdf_to_docx",
        staticmethod(lambda *a, **k: called.setdefault("pdf2docx", True) or None),
    )
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=[], bounding_boxes=[],
        config=RedactionConfig(replacement_mode="pseudonym"),
    )
    assert "pdf2docx" not in called, "零实体不应触发转换链路"
    assert result["redacted_count"] == 0
    with open(src, "rb") as f1, open(result["output_path"], "rb") as f2:
        assert f1.read() == f2.read(), "零实体产物应与原文件一致"


@pytest.mark.asyncio
async def test_pdf_replacement_per_entity_check_catches_dropped(_dirs, monkeypatch):
    """逐实体校验（评审 I1）：fake pdf2docx 丢掉实体 B（源 docx 就没有 B），
    即使 A 全部替换成功也必须回退原位替换——聚合计数会被『A 多次+B 零次』
    凑数骗过。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    called = {}

    def _fake_pdf2docx(src_path, wd):
        from docx import Document as _Doc
        fake_docx = os.path.join(wd, "source.docx")
        d = _Doc()
        d.add_paragraph("委托人：陈明飞")  # 只保留 A，B 被「转换丢失」
        d.save(fake_docx)
        called["pdf2docx"] = True
        return fake_docx

    async def _fake_docx2pdf(docx, out):
        called["docx2pdf"] = True
        import shutil; shutil.copy(str(src), out); return True

    monkeypatch.setattr(Redactor, "_pdf_to_docx", staticmethod(_fake_pdf2docx))
    monkeypatch.setattr(Redactor, "_docx_to_pdf", staticmethod(_fake_docx2pdf))
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=_entities(), bounding_boxes=[],
        config=RedactionConfig(replacement_mode="structured"),
    )
    assert called.get("pdf2docx")
    assert "docx2pdf" not in called, "检测到转换丢实体后必须回退，不得交付回转产物"
    text = "".join(p.get_text() for p in fitz.open(result["output_path"]))
    squeezed = text.replace(" ", "")
    assert "陈明飞" not in squeezed and "110101199003078515" not in squeezed, \
        "回退原位替换后两个实体原文都必须消失"


@pytest.mark.asyncio
async def test_pdf_replacement_preserved_entity_not_flagged_residual(_dirs, monkeypatch):
    """设计性保留实体豁免残留判定（评审跟进项）：公共机构替换词==原文
    （org_rules #56），不得因此永远回退原位替换、废掉 docx 回转收益。"""
    up, _ = _dirs
    src = up / "t.pdf"
    _make_pdf(src)
    called = {}

    def _fake_pdf2docx(src_path, wd):
        from docx import Document as _Doc
        fake_docx = os.path.join(wd, "source.docx")
        d = _Doc()
        d.add_paragraph("主管机关：司法部")
        d.save(fake_docx)
        called["pdf2docx"] = True
        return fake_docx

    async def _fake_docx2pdf(docx, out):
        called["docx2pdf"] = True
        import shutil; shutil.copy(str(src), out); return True

    monkeypatch.setattr(Redactor, "_pdf_to_docx", staticmethod(_fake_pdf2docx))
    monkeypatch.setattr(Redactor, "_docx_to_pdf", staticmethod(_fake_docx2pdf))
    ents = [Entity(id="e1", text="司法部", type="INSTITUTION_NAME", start=0, end=3, page=1, selected=True)]
    result = await Redactor().redact(
        file_info={"file_path": str(src), "file_type": "pdf"},
        entities=ents, bounding_boxes=[],
        config=RedactionConfig(
            replacement_mode="pseudonym",
            custom_replacements={"司法部": "司法部"},  # 保留原文
        ),
    )
    assert called.get("pdf2docx") and called.get("docx2pdf"), \
        "保留实体（替换词==原文）不应触发残留回退，应正常走 docx 回转"
    assert result["residual_entities"] == []


def test_soffice_staging_root_snap_vs_nonsnap():
    """snap 版 soffice 暂存走 $HOME 非隐藏目录（沙箱读不了 /tmp 和 ~/.cache，
    评审 C2）；非 snap 走系统 /tmp。"""
    from app.services.redaction.text_redactor import TextRedactorMixin

    snap_root = TextRedactorMixin._soffice_staging_root("/snap/bin/libreoffice")
    home = os.path.expanduser("~")
    assert snap_root.startswith(home), "snap 暂存必须在 $HOME 下"
    assert not os.path.basename(snap_root).startswith("."), "snap 暂存不能是隐藏目录"
    import tempfile as _tmp
    assert TextRedactorMixin._soffice_staging_root("/usr/bin/soffice") == os.path.join(
        _tmp.gettempdir(), "redaction-soffice"
    )


def test_locate_merges_fragmented_rects_and_handles_spaces(_dirs):
    """定位核心空白鲁棒（#66 验收反馈）：WPS/LibreOffice 文本层的怪空格
    会让 search_for 返回逐字块碎矩形（一个日期 6 框）或 0 命中——
    ①行内合并成整框；②双向空白归一化兜底。"""
    up, _ = _dirs
    src = up / "sp.pdf"
    doc = fitz.open()
    page = doc.new_page()
    # 模拟 WPS 产物：日期带怪空格
    page.insert_text((72, 130), "签署日期：2022 年1 月25 日 由双方确认", fontsize=12, fontname="china-s")
    doc.save(str(src))
    doc.close()

    from app.services.redactor import Redactor

    # ① NER 返回带空格形态（与页面一致）→ 原样命中，碎矩形行内合并为 1 框
    boxes, missed = Redactor.locate_entity_texts(
        str(src), [("2022 年1 月25 日", "DATE")], "ner", "ner"
    )
    assert not missed
    assert len(boxes) == 1, f"碎矩形应行内合并为 1 框，实际 {len(boxes)}"

    # ② NER 返回无空格形态（与页面不一致）→ 字符映射兜底仍命中
    boxes2, missed2 = Redactor.locate_entity_texts(
        str(src), [("2022年1月25日", "DATE")], "ner", "ner"
    )
    assert not missed2, "无空格形态应经字符映射兜底命中"
    assert len(boxes2) == 1

    # ③ source 透传
    assert boxes[0].source == "ner"


def test_locate_numeric_anchor_for_rewritten_case_number(_dirs):
    """数字锚点兜底（#66 复验反馈）：NER 改写案号文本（"英检"→"英德"、
    括号全半角规范化）后字面搜索永远失败——用唯一长数字串锚定整行。"""
    up, _ = _dirs
    src = up / "case.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 130), "案件编号（审查起诉）：英检刑诉受[2022]441881000050 号", fontsize=12, fontname="china-s")
    doc.save(str(src))
    doc.close()

    from app.services.redactor import Redactor

    # NER 改写形态：全角六角括号 + "英德"补全——字面无命中，数字锚点兜底
    boxes, missed = Redactor.locate_entity_texts(
        str(src), [("英德刑诉受〔2022〕441881000050 号", "CASE_NUMBER")], "ner", "ner"
    )
    assert not missed, "案号应经数字锚点兜底命中"
    assert len(boxes) == 1
    assert boxes[0].source == "ner"

    # 多个/无长数字串的文本不走锚点（防误匹配）
    boxes2, missed2 = Redactor.locate_entity_texts(
        str(src), [("2022 年01 月25 日受理", "DATE")], "ner", "ner"
    )
    assert "2022 年01 月25 日受理" in missed2, "多数字串文本不得用锚点乱框"


def test_locate_finds_all_occurrences_across_forms(_dirs):
    """同页多形态不短路（#66 复验反馈）：同一日期同页出现两种空格形态
    （收案行紧排 + 正文跨行），必须全部框出——原样命中一处就停会漏掉
    其余出现位置。"""
    up, _ = _dirs
    src = up / "multi.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "收案时间：2022 年1 月25 日", fontsize=12, fontname="china-s")
    page.insert_text((72, 200), "我院于2022 年01 月25 日受理该案", fontsize=12, fontname="china-s")
    # 正文：跨行形态（行尾 2022 年1 月 / 行首 25 日）
    page.insert_text((72, 300), "于2022 年1 月", fontsize=12, fontname="china-s")
    page.insert_text((72, 330), "25 日告知当事人", fontsize=12, fontname="china-s")
    doc.save(str(src))
    doc.close()

    from app.services.redactor import Redactor

    boxes, missed = Redactor.locate_entity_texts(
        str(src), [("2022 年1 月25 日", "DATE")], "ner", "ner"
    )
    assert not missed
    assert len(boxes) >= 2, f"同页多处出现应全部框出（含跨行），实际 {len(boxes)} 框"
    # 收案行（y≈100）必须有框
    assert any(b.y < 0.5 for b in boxes), "第一处（收案行）应有框"
