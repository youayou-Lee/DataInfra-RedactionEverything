"""合成 docx 生成器（Issue #37 评测集）：语料行 → docx 文件。

docx 无真实分页概念，GT 聚合为单条「第 1 页」（run_eval 对 docx 做文档级实体对齐）。
"""

from __future__ import annotations

import io
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gen_text  # noqa: E402

_FIXED_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)  # 固定 core.xml 时间戳，保证产物确定性
_FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)  # zip 条目时间戳（默认取当前时间，破坏字节确定性）


def _freeze_zip(source: Path, out_path: Path) -> None:
    """重写 zip 条目时间戳为固定值（python-docx 保存时写当前时间）。"""
    with zipfile.ZipFile(str(source)) as zf_in, \
            zipfile.ZipFile(str(out_path), "w", zipfile.ZIP_DEFLATED) as zf_out:
        for info in zf_in.infolist():
            new_info = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_DATE)
            new_info.compress_type = info.compress_type
            new_info.external_attr = info.external_attr
            new_info.create_system = info.create_system
            zf_out.writestr(new_info, zf_in.read(info.filename))


def build_docx(out_path: Path, *, pages: int, doc_type: str = "contract",
               density: str = "mid") -> list[dict]:
    title = gen_text.DOC_TITLES[doc_type]
    document = Document()
    document.core_properties.created = _FIXED_TIME
    document.core_properties.modified = _FIXED_TIME
    document.core_properties.last_modified_by = "eval-gen"
    entities_all: dict[str, list[str]] = {}
    for i in range(pages):
        data = gen_text.build_page(i, doc_type=doc_type, density=density)
        document.add_heading(f"{title}（第 {i + 1} 段 / 共 {pages} 段）", level=2)
        for line in data["lines"]:
            document.add_paragraph(line)
        for etype, values in data["entities"].items():
            entities_all.setdefault(etype, []).extend(values)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_suffix(out_path.suffix + ".tmp")
    document.save(str(tmp))
    _freeze_zip(tmp, out_path)
    tmp.unlink(missing_ok=True)
    merged = {k: sorted(set(v)) for k, v in entities_all.items() if v}
    return [{"page": 0, "entities": merged}]
