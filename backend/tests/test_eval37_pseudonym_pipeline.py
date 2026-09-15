"""Issue #37 化名管线纯函数单测：CSV 往返、重复化名拒绝、GT 定位与 missing 检测。"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "eval" / "scripts"


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


bps = _load("build_pseudonym_set_under_test", SCRIPTS_DIR / "build_pseudonym_set.py")

ROWS = [
    {"原文": "王建国", "类型": "姓名", "化名": "王某1"},
    {"原文": "11010119600101000X", "类型": "身份证号", "化名": "110101196001010011"},
    {"原文": "华宸信息技术有限公司", "类型": "机构名称", "化名": "某公司1"},
]


def test_mapping_csv_roundtrip(tmp_path):
    csv_path = tmp_path / "reviewed.csv"
    bps.write_mapping_csv(csv_path, ROWS)
    back = bps.read_mapping_csv(csv_path)
    assert [r["化名"] for r in back] == [r["化名"] for r in ROWS]


def test_read_mapping_rejects_duplicate_pseudonym(tmp_path):
    csv_path = tmp_path / "dup.csv"
    bps.write_mapping_csv(csv_path, ROWS + [{"原文": "李四", "类型": "姓名", "化名": "王某1"}])
    with pytest.raises(ValueError, match="重复化名"):
        bps.read_mapping_csv(csv_path)


def test_build_merged_entities_includes_manual_rows():
    """I1：补行（原文不在识别结果中）并入执行载荷，附 manual 标记。"""
    entities = [{"id": "eval37-1", "text": "王建国", "type": "PERSON", "start": 0, "end": 3, "page": 1}]
    rows = [{"原文": "王建国", "类型": "姓名", "化名": "王某1"},
            {"原文": "李四", "类型": "姓名", "化名": "李某1"}]  # 补行：漏检实体
    merged = bps.build_merged_entities(entities, rows)
    assert len(merged) == 2
    manual = next(e for e in merged if e.get("manual"))
    assert manual["text"] == "李四" and manual["type"] == "PERSON"  # 中文名反查英文 ID


def test_read_mapping_skips_empty_rows(tmp_path):
    csv_path = tmp_path / "holes.csv"
    bps.write_mapping_csv(csv_path, ROWS + [{"原文": "", "类型": "", "化名": ""}])
    assert len(bps.read_mapping_csv(csv_path)) == 3


def test_locate_gt_txt(tmp_path):
    target = tmp_path / "p.txt"
    target.write_text("甲方王某1 与某公司1 签订合同。\n见证人：王某1。", encoding="utf-8")
    located = bps.locate_gt(target, ROWS)
    assert located["missing"] == [{"原文": "11010119600101000X", "类型": "身份证号",
                                   "化名": "110101196001010011"}]
    entities = located["pages"][0]["entities"]
    assert entities["姓名"] == ["王某1"]
    assert entities["机构名称"] == ["某公司1"]
    assert "身份证号" not in entities


def test_locate_gt_docx_pages_aggregated(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("甲方王某1。")
    document.add_paragraph("某公司1 盖章。")
    target = tmp_path / "p.docx"
    document.save(str(target))
    located = bps.locate_gt(target, ROWS)
    assert located["page_count"] == 1  # docx 文档级聚合
    entities = located["pages"][0]["entities"]
    assert entities["姓名"] == ["王某1"] and entities["机构名称"] == ["某公司1"]


def test_locate_gt_pdf_page_granularity(tmp_path):
    import fitz
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "第一页：王某1 签字", fontname="china-s")
    doc.new_page().insert_text((72, 72), "第二页：某公司1 盖章", fontname="china-s")
    target = tmp_path / "p.pdf"
    doc.save(str(target))
    doc.close()
    located = bps.locate_gt(target, ROWS)
    assert located["page_count"] == 2
    assert located["pages"][0]["entities"] == {"姓名": ["王某1"]}
    assert located["pages"][1]["entities"] == {"机构名称": ["某公司1"]}
