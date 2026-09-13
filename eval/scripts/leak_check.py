"""入库前残留检测（Issue #37，设计文档 D5）：化名版成品中不得出现对照表原文。

双保险：
  ①（本地，必跑）对照表每个原文串对化名版全文做归一化 grep；
  ②（--api-base 时）对化名版跑一次识别，识别结果 ∩ 原文串 也算残留。

用法：
  python eval/scripts/leak_check.py eval/datasets/pseudonymized/xxx.pdf \
      --mapping /path/to/private/mapping.csv            # 私有目录，不入库
  python eval/scripts/leak_check.py xxx.docx --mapping m.csv --api-base http://127.0.0.1:8000

对照表 CSV 列：原文,类型,化名（utf-8-sig，T1 对照表同格式）。
退出码：0 = 零残留可入库；1 = 发现残留（明细见 stdout / --out JSON）。
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common_api  # noqa: E402

_FT_RE = re.compile(r"[，。、；：？！“”‘’（）【】《》\s,.;:?!\"'()\[\]<>\-—_·]")


def _norm(text: str) -> str:
    """强归一：去空白 + 去常见标点 + 连字符/下划线 + 大小写折叠。

    覆盖面不得弱于评测口径 nerq._normalize（去 [\\s\\-—_] + upper）——检测闸门比指标
    更宽松等于放行（评审 I3）。
    """
    return _FT_RE.sub("", str(text or "")).upper()


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import fitz
        with fitz.open(str(path)) as doc:
            return "\n".join(page.get_text() for page in doc)
    if suffix == ".docx":
        from docx import Document
        document = Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    raise ValueError(f"不支持的文件类型: {path}")


def load_mapping(csv_path: Path) -> list[dict]:
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    required = {"原文", "类型", "化名"}
    if not rows or not required.issubset(rows[0].keys()):
        raise ValueError(f"对照表缺少列 {required - set(rows[0].keys() if rows else [])}")
    return rows


def run_check(target: Path, mapping_rows: list[dict]) -> dict:
    text = extract_text(target)
    text_norm = _norm(text)
    findings = []
    for row in mapping_rows:
        raw = (row.get("原文") or "").strip()
        if not raw:
            continue
        if _norm(raw) and _norm(raw) in text_norm:
            findings.append({"原文": raw, "类型": row.get("类型", ""), "how": "grep"})
    return {"file": str(target), "checked_originals": len(mapping_rows), "findings": findings,
            "clean": not findings, "rescan": None, "text_chars": len(text)}


def run_rescan(api: common_api.EvalApi, target: Path, mapping_rows: list[dict]) -> dict:
    """保险②：识别化名版，预测实体与原文串求交（归一域）。

    pdf 走逐页 vision；docx/txt 走 parse + hybrid NER（D7：vision 不支持该类载体）。
    """
    file_id = api.upload(target)
    try:
        suffix = target.suffix.lower()
        page_entities: list[dict[str, list[str]]] = []
        if suffix == ".pdf":
            import fitz
            with fitz.open(str(target)) as doc:
                pages = doc.page_count
            for page in range(1, pages + 1):
                page_entities.append(common_api.extract_page_entities(api.vision(file_id, page)))
        elif suffix in (".docx", ".txt"):
            entities, _ = api.parse_and_hybrid_ner(file_id)
            page_entities.append(entities)
            pages = 1
        else:
            raise ValueError(f"不支持: {target}")
        findings = []
        originals = {(row.get("原文") or "").strip(): row for row in mapping_rows
                     if (row.get("原文") or "").strip()}
        for page_no, entities in enumerate(page_entities, start=1):
            for etype, values in entities.items():
                for value in values:
                    hit = originals.get(value) or next(
                        (o for o in originals if _norm(o) == _norm(value)), None)
                    if hit:
                        findings.append({"原文": hit["原文"], "类型": etype, "how": "rescan",
                                         "page": page_no, "识别为": value})
        return {"pages": pages, "findings": findings, "clean": not findings}
    finally:
        api.delete_file(file_id)


def main() -> int:
    parser = argparse.ArgumentParser(description="化名版入库前残留检测（Issue #37 D5）")
    parser.add_argument("target", type=Path, help="化名版成品文件（pdf/docx/txt）")
    parser.add_argument("--mapping", required=True, type=Path, help="对照表 CSV（私有目录，不入库）")
    parser.add_argument("--api-base", default=None, help="可选：识别复扫的 backend 地址")
    parser.add_argument("--api-user", default="eval_user")
    parser.add_argument("--api-pass", default="EvalUser!2026")
    parser.add_argument("--out", default=None, type=Path, help="结果 JSON 输出路径")
    args = parser.parse_args()

    rows = load_mapping(args.mapping)
    result = run_check(args.target, rows)
    if args.api_base:
        api = common_api.EvalApi(args.api_base, args.api_user, args.api_pass)
        try:
            result["rescan"] = run_rescan(api, args.target, rows)
        finally:
            api.close()
        result["clean"] = result["clean"] and result["rescan"]["clean"]

    print(json.dumps(result, ensure_ascii=False, indent=2))
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if not result["clean"]:
        print(f"❌ 发现 {len(result['findings']) + len((result['rescan'] or {}).get('findings', []))} 处残留，禁止入库",
              file=sys.stderr)
        return 1
    print("✅ 零残留，可入库")
    return 0


if __name__ == "__main__":
    sys.exit(main())
